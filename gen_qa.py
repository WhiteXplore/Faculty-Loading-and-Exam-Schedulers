from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = Workbook()

COLOURS = {
    "Admin":   {"header": "1F4E79", "group": "BDD7EE", "row_even": "DEEAF1"},
    "PC":      {"header": "375623", "group": "C6EFCE", "row_even": "E2EFDA"},
    "Faculty": {"header": "7B2C2C", "group": "F4CCCC", "row_even": "FCE4D6"},
    "Shared":  {"header": "3D3D3D", "group": "D9D9D9", "row_even": "F2F2F2"},
    "Integ":   {"header": "4B0082", "group": "E1D5F5", "row_even": "F3EEFE"},
}

thin = Side(style="thin", color="AAAAAA")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

def fill(hex_colour):
    return PatternFill("solid", fgColor=hex_colour)

def write_header(ws, columns, col_widths, header_hex):
    ws.append(columns)
    for idx, (col, width) in enumerate(zip(columns, col_widths), 1):
        cell = ws.cell(row=1, column=idx)
        cell.font = Font(bold=True, color="FFFFFF", name="Calibri", size=10)
        cell.fill = fill(header_hex)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.row_dimensions[1].height = 28

def write_group_row(ws, label, num_cols, group_hex):
    ws.append([label] + [""] * (num_cols - 1))
    row = ws.max_row
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=num_cols)
    cell = ws.cell(row=row, column=1)
    cell.font = Font(bold=True, name="Calibri", size=10)
    cell.fill = fill(group_hex)
    cell.alignment = Alignment(horizontal="left", vertical="center")
    cell.border = border
    ws.row_dimensions[row].height = 18

def write_row(ws, data, even, palette_key):
    ws.append(data)
    row = ws.max_row
    bg = COLOURS[palette_key]["row_even"] if even else "FFFFFF"
    for idx in range(1, len(data) + 1):
        cell = ws.cell(row=row, column=idx)
        cell.fill = fill(bg)
        cell.font = Font(name="Calibri", size=10)
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        cell.border = border
    ws.row_dimensions[row].height = 42

# ── Sheet 1: Shared Auth ─────────────────────────────────────────
ws = wb.active
ws.title = "Shared - Auth"
ws.freeze_panes = "A2"
cols   = ["Test ID", "User Type", "Test Case", "Steps", "Expected Result", "Status", "Remarks"]
widths = [10, 20, 30, 48, 40, 12, 20]
write_header(ws, cols, widths, COLOURS["Shared"]["header"])

auth_rows = [
    ("A-01","All Roles","Valid login - Admin","Enter valid Admin email & password > click Login","Redirected to /admin-dashboard","",""),
    ("A-02","All Roles","Valid login - Program Chairperson","Enter valid PC credentials > click Login","Redirected to /progchair-dashboard","",""),
    ("A-03","All Roles","Valid login - Faculty","Enter valid Faculty credentials > click Login","Redirected to /faculty-dashboard","",""),
    ("A-04","All Roles","Invalid login - wrong password","Enter correct email + wrong password > Login","Error message shown; no redirect","",""),
    ("A-05","All Roles","Invalid login - wrong email","Enter non-existent email > Login","Error message shown; no redirect","",""),
    ("A-06","All Roles","Empty login form","Click Login with all fields blank","Field validation errors shown","",""),
    ("A-07","Faculty","Unauthorized page - Admin route","Login as Faculty > manually navigate to /admin-dashboard","Redirected to 404 Not Found","",""),
    ("A-08","Program Chairperson","Unauthorized page - Admin route","Login as PC > manually navigate to /instructors","Redirected to 404 Not Found","",""),
    ("A-09","All Roles","Logout","Click Logout from any role","Session cleared; redirected to login page","",""),
    ("A-10","All Roles","Profile View","Navigate to /profile-view from any role","Own profile data displayed correctly","",""),
]
for i, r in enumerate(auth_rows):
    write_row(ws, list(r), i % 2 == 0, "Shared")

# ── Sheet 2: Admin ───────────────────────────────────────────────
ws2 = wb.create_sheet("Admin")
ws2.freeze_panes = "A2"
cols2   = ["Test ID", "Feature Group", "Test Case", "Steps", "Expected Result", "Status", "Remarks"]
widths2 = [10, 24, 32, 50, 42, 12, 20]
write_header(ws2, cols2, widths2, COLOURS["Admin"]["header"])

admin_data = [
    ("Institutes", [
        ("AD-01","View institutes list","Navigate to /institutes","Institute list loads correctly"),
        ("AD-02","Add new institute","Click Add > fill name & code > Submit","New institute appears in list"),
        ("AD-03","Edit institute","Click Edit > change name > Save","Updated name reflected in list"),
        ("AD-04","Delete institute","Click Delete > Confirm","Institute removed from list"),
        ("AD-05","Add duplicate institute","Add institute with same name/code as existing","Error or warning shown; no duplicate created"),
    ]),
    ("Curriculum", [
        ("AD-06","View all curricula","Navigate to /curriculum","List loads with program and institute data"),
        ("AD-07","Add curriculum","Select institute, program, start year, end year > Submit","Curriculum appears in list"),
        ("AD-08","Add duplicate curriculum","Same institute + program + start/end year","Returns existing record; no duplicate created"),
        ("AD-09","Edit curriculum years","Change start or end year > Save","Updated values shown"),
        ("AD-10","Delete curriculum","Click Delete > Confirm","Curriculum removed from list"),
    ]),
    ("Curriculum Courses", [
        ("AD-11","View courses linked to curriculum","Open curriculum courses page","Linked courses listed correctly"),
        ("AD-12","Add course to curriculum","Select curriculum + course > Submit","Course appears under curriculum"),
        ("AD-13","Add duplicate curriculum-course","Same curriculum + same course","Duplicate prevented or warning shown"),
        ("AD-14","Remove course from curriculum","Click Delete on linked course > Confirm","Course unlinked from curriculum"),
    ]),
    ("Courses", [
        ("AD-15","View all courses","Navigate to /courses","Full course list loads"),
        ("AD-16","Add course","Fill code, title, lec, lab, level, semester > Submit","Course added to list"),
        ("AD-17","Edit course","Change title or units > Save","Changes reflected"),
        ("AD-18","Delete course","Click Delete > Confirm","Course removed"),
        ("AD-19","View curriculum offer report","Navigate to /report-curriculum-offers","Report grouped by institute loads correctly"),
        ("AD-20","View offers by institute","Click /view-curriculum-offers/:institute_id","Courses per institute shown via curriculum_courses"),
    ]),
    ("Rooms", [
        ("AD-21","View rooms","Navigate to /rooms","Room list loads"),
        ("AD-22","Add room","Fill room details > Submit","Room appears in list"),
        ("AD-23","Edit room","Change name or capacity > Save","Changes reflected"),
        ("AD-24","Delete room","Click Delete > Confirm","Room removed from list"),
    ]),
    ("School Years", [
        ("AD-25","View school years","Navigate to /school-years","School years list loads"),
        ("AD-26","Add school year","Enter start/end year > Submit","New school year created and shown"),
        ("AD-27","Edit school year","Update data > Save","Changes saved"),
        ("AD-28","Delete school year","Click Delete > Confirm","Removed from list"),
        ("AD-29","Active school year on dashboard","Load admin dashboard","Correct latest active school year displayed"),
    ]),
    ("Faculty Management", [
        ("AD-30","View all instructors","Navigate to /instructors","Full faculty list loads"),
        ("AD-31","Add instructor manually","Fill name, employment type, unit load, role > Submit","Instructor appears in list"),
        ("AD-32","Import instructors via Excel","Upload .xlsx file > Submit","Instructors imported successfully"),
        ("AD-33","Import faculty expertise via Excel","Upload expertise .xlsx > Submit","Expertise data linked to faculty"),
        ("AD-34","Edit instructor","Change employment type or unit load > Save","Updated"),
        ("AD-35","Delete instructor","Click Delete > Confirm","Removed"),
    ]),
    ("Year & Section (Classes)", [
        ("AD-36","View class sections","Navigate to /year-section","Class sections list loads"),
        ("AD-37","Add class section","Fill set name, program, class size, school year > Submit","Section created"),
        ("AD-38","Edit class section","Change class size or set name > Save","Changes reflected"),
        ("AD-39","Delete class section","Click Delete > Confirm","Section removed"),
    ]),
    ("Faculty Loading & Schedule", [
        ("AD-40","View faculty loads overview","Navigate to /faculty-loads","Overview loads with assigned faculty"),
        ("AD-41","Assign class to faculty","Select faculty + class > Assign","Assignment saved"),
        ("AD-42","Run GA scheduler","Trigger schedule generation","Python GA runs; JSON output produced"),
        ("AD-43","Read generated JSON schedule","Open generated schedule view","Output displayed correctly on screen"),
        ("AD-44","Save final schedule (bulk)","Save generated schedule > Submit","Records saved to final_generated_class_schedules"),
        ("AD-45","View saved final schedules","Navigate to /view-faculty-loads","Final schedules visible"),
    ]),
    ("User Account Management", [
        ("AD-46","View all user accounts","Navigate to /user-accounts","Full user list loads"),
        ("AD-47","Add user account","Fill name, email, role, program > Submit","User created"),
        ("AD-48","Edit user account","Change role or program > Save","Changes saved"),
        ("AD-49","Delete user account","Click Delete > Confirm","User removed"),
    ]),
]

even = False
for group_label, group_rows in admin_data:
    write_group_row(ws2, "  " + group_label, len(cols2), COLOURS["Admin"]["group"])
    for r in group_rows:
        write_row(ws2, [r[0], group_label, r[1], r[2], r[3], "", ""], even, "Admin")
        even = not even

# ── Sheet 3: Program Chairperson ─────────────────────────────────
ws3 = wb.create_sheet("Program Chairperson")
ws3.freeze_panes = "A2"
write_header(ws3, cols2, widths2, COLOURS["PC"]["header"])

pc_data = [
    ("Dashboard", [
        ("PC-01","View dashboard","Login as Program Chairperson","progchair-dashboard loads with program data"),
    ]),
    ("Courses (Program-scoped)", [
        ("PC-02","View courses for own program","Navigate to /program-chair-courses","Only program-relevant courses shown"),
        ("PC-03","Add course","Fill details > Submit","Course added"),
        ("PC-04","Edit course","Modify course > Save","Changes reflected"),
    ]),
    ("Year & Section / Assigned Classes", [
        ("PC-05","View year sections","Navigate to /program-chair-year-section","Sections for own program shown"),
        ("PC-06","Sync assigned courses from classes","Trigger sync","program_year_courses populated correctly"),
        ("PC-07","View class assigned classes overview","Navigate to /program-chair-class-assigned-classes","Courses per class visible"),
    ]),
    ("Faculty (Program-scoped)", [
        ("PC-08","View faculty list","Navigate to /program-chair-faculty-list","Only faculty under own program shown"),
        ("PC-09","View faculty expertise","Navigate to /faculty-expertise","Expertise per faculty listed"),
    ]),
    ("Faculty Loading", [
        ("PC-10","View faculty loads","Navigate to /program-chair-faculty-loads","Loads for own program shown"),
        ("PC-11","View program faculty loading","Navigate to /program-faculty-loading","Detailed per-faculty load visible"),
    ]),
    ("Final Schedules", [
        ("PC-12","View final schedules","Navigate to /program-final-schedules","Generated schedules for program shown"),
    ]),
    ("Programs", [
        ("PC-13","View programs list","Navigate to /program-programs","Programs list visible"),
    ]),
    ("Profile", [
        ("PC-14","View / edit profile","Navigate to /profile-view","Own profile shown and editable"),
    ]),
]

even = False
for group_label, group_rows in pc_data:
    write_group_row(ws3, "  " + group_label, len(cols2), COLOURS["PC"]["group"])
    for r in group_rows:
        write_row(ws3, [r[0], group_label, r[1], r[2], r[3], "", ""], even, "PC")
        even = not even

# ── Sheet 4: Faculty ─────────────────────────────────────────────
ws4 = wb.create_sheet("Faculty")
ws4.freeze_panes = "A2"
write_header(ws4, cols2, widths2, COLOURS["Faculty"]["header"])

faculty_data = [
    ("Dashboard", [
        ("FA-01","View dashboard","Login as Faculty","faculty-dashboard loads"),
    ]),
    ("My Loading", [
        ("FA-02","View assigned courses/load","Navigate to /faculty-load","Assigned courses and schedule shown"),
        ("FA-03","Faculty with no assignments","Login as faculty with no assigned classes","Empty state shown; no error"),
    ]),
    ("My Preference / Expertise", [
        ("FA-04","View current preferences","Navigate to /faculty-preference","Expertise and preferred time visible"),
        ("FA-05","Set preferred time","Select time slot > Save","Preference saved and reflected"),
        ("FA-06","Update expertise","Add or remove course expertise > Save","Expertise updated correctly"),
    ]),
    ("Profile", [
        ("FA-07","View / edit profile","Navigate to /profile-view","Own data shown and editable"),
    ]),
    ("Access Restriction", [
        ("FA-08","Access Admin page","Manually navigate to /instructors","Redirected to 404 Not Found"),
        ("FA-09","Access PC page","Manually navigate to /progchair-dashboard","Redirected to 404 Not Found"),
    ]),
]

even = False
for group_label, group_rows in faculty_data:
    write_group_row(ws4, "  " + group_label, len(cols2), COLOURS["Faculty"]["group"])
    for r in group_rows:
        write_row(ws4, [r[0], group_label, r[1], r[2], r[3], "", ""], even, "Faculty")
        even = not even

# ── Sheet 5: Integration ─────────────────────────────────────────
ws5 = wb.create_sheet("Integration")
ws5.freeze_panes = "A2"
cols5   = ["Test ID", "Test Case", "Steps", "Expected Result", "Status", "Remarks"]
widths5 = [10, 36, 58, 46, 12, 20]
write_header(ws5, cols5, widths5, COLOURS["Integ"]["header"])

integ_rows = [
    ("INT-01","Full faculty loading flow",
     "Admin: create school year > add classes > link courses via curriculum_courses > assign to faculty > run GA > save final schedule",
     "Schedule saved and visible in final schedules"),
    ("INT-02","Curriculum offer report integrity",
     "Admin links courses to curriculum via curriculum_courses > open report",
     "Report shows correct courses grouped by institute using junction table"),
    ("INT-03","Faculty preference reflected in GA output",
     "Faculty sets preferred time > Admin runs GA scheduler",
     "GA output respects faculty preferred time"),
    ("INT-04","Program Chairperson sees correct program scope",
     "Login as PC > check courses, faculty loads, and sections",
     "All views only show data for the PC program_id"),
    ("INT-05","Calendar events visible across roles",
     "Admin adds a calendar event > PC and Faculty view calendar",
     "Event visible across all authenticated roles"),
    ("INT-06","Sync program year courses integrity",
     "PC triggers sync > check program_year_courses table",
     "Records match classes + curriculum_courses junction correctly"),
]

for i, r in enumerate(integ_rows):
    write_row(ws5, list(r) + ["", ""], i % 2 == 0, "Integ")

# ── Save ─────────────────────────────────────────────────────────
out = r"C:\Users\JiMkErR\Documents\DNSCSYSTEM\GA-faculty-scheduler\GA-Faculty-Loading-Final\Faculty-Loading-and-Exam-Schedulers\QA_TestCases_FacultyLoader.xlsx"
wb.save(out)
print("Saved:", out)
