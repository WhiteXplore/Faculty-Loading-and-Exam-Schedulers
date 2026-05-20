from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

PATH = r"C:\Users\JiMkErR\Documents\DNSCSYSTEM\GA-faculty-scheduler\GA-Faculty-Loading-Final\Faculty-Loading-and-Exam-Schedulers\QA_TestCases_FacultyLoader.xlsx"
wb = load_workbook(PATH)
ws = wb["Integration"]

thin = Side(style="thin", color="AAAAAA")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

def fill(h):
    return PatternFill("solid", fgColor=h)

HEADER_HEX = "722F37"
GROUP_HEX  = "F2DCDB"
ROW_EVEN   = "FDF2F0"
N = 6

def divider(ws, label):
    ws.append([label] + [""] * (N - 1))
    r = ws.max_row
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=N)
    c = ws.cell(r, 1)
    c.font = Font(bold=True, name="Calibri", size=11, color="FFFFFF")
    c.fill = fill(HEADER_HEX)
    c.alignment = Alignment(horizontal="left", vertical="center")
    c.border = border
    ws.row_dimensions[r].height = 22

def group(ws, label):
    ws.append([label] + [""] * (N - 1))
    r = ws.max_row
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=N)
    c = ws.cell(r, 1)
    c.font = Font(bold=True, name="Calibri", size=10)
    c.fill = fill(GROUP_HEX)
    c.alignment = Alignment(horizontal="left", vertical="center")
    c.border = border
    ws.row_dimensions[r].height = 18

def add_row(ws, data, even):
    ws.append(data)
    r = ws.max_row
    bg = ROW_EVEN if even else "FFFFFF"
    for i in range(1, N + 1):
        c = ws.cell(r, i)
        c.fill = fill(bg)
        c.font = Font(name="Calibri", size=10)
        c.alignment = Alignment(vertical="top", wrap_text=True)
        c.border = border
    ws.row_dimensions[r].height = 55

ws.append([""] * N)
divider(ws, "  SCHEDULING CONSTRAINT TEST CASES  (from GA Scheduler rules)")

data = [
    ("Hard Constraint - No Overlap", [
        ("CON-01", "No faculty schedule overlap",
         "Assign the same faculty member to two different classes at overlapping times on the same day, then run the GA scheduler.",
         "The GA must not produce a schedule where one faculty member is teaching two classes at the same time on the same day."),
        ("CON-02", "No room double-booking",
         "Assign two different classes to the same room at overlapping times on the same day, then run the GA scheduler.",
         "The GA must not schedule two classes in the same room at the same time on the same day."),
        ("CON-03", "No class section overlap",
         "Set up one class section (e.g., 1st Year BSCS) with two courses at overlapping times, then run the GA scheduler.",
         "The same class section must not have two courses running at the same time on the same day."),
    ]),
    ("Hard Constraint - Lunch Break", [
        ("CON-04", "No class scheduled during lunch break (12:00 PM - 1:00 PM)",
         "Run the GA scheduler and inspect all generated time slots across the output.",
         "No class must start or overlap with the 12:00 PM to 1:00 PM window. All schedules must skip this hour entirely."),
    ]),
    ("Hard Constraint - Part-Time Faculty Hours", [
        ("CON-05", "Part-time faculty cannot teach before 5:00 PM",
         "Tag a faculty member as part-time, run the GA, then inspect all time slots assigned to that faculty.",
         "All classes assigned to a part-time faculty must start at 5:00 PM (17:00) or later. No earlier slot is allowed."),
        ("CON-06", "Part-time faculty cannot teach beyond 10:00 PM",
         "Tag a faculty member as part-time, run the GA, then check whether any of their classes end after 10:00 PM.",
         "No class assigned to a part-time faculty should extend past 10:00 PM (22:00)."),
    ]),
    ("Hard Constraint - Room Capacity", [
        ("CON-07", "Room capacity exceeded beyond 5-student tolerance is rejected",
         "Assign a class of 61 students to a room with capacity 55 (6 over the limit), then run the GA.",
         "The GA must not assign that class to the room. An excess of more than 5 students over room capacity is not allowed."),
        ("CON-08", "Room capacity within 5-student tolerance is accepted",
         "Assign a class of 59 students to a room with capacity 55 (4 students over, within tolerance), then run the GA.",
         "The GA may assign this class to the room since the excess is within the allowed 5-student tolerance."),
    ]),
    ("Hard Constraint - Faculty Expertise", [
        ("CON-09", "Faculty can only be assigned courses matching their expertise",
         "Create a faculty with no expertise entry for a specific course, run the GA, and check if that course is assigned to them.",
         "The course must not be assigned to that faculty. Only faculty with a matching expertise entry may receive the course."),
        ("CON-10", "Unassigned course logged when no qualified faculty exists",
         "Set up a course for which no faculty has a matching expertise entry, then run the GA.",
         "The course must appear in the unscheduled meetings output with the reason stating no qualified faculty was found."),
    ]),
    ("Hard Constraint - Faculty Unit Load", [
        ("CON-11", "Faculty unit load must not be exceeded",
         "Set a faculty unit_load limit (e.g., 21 units), assign courses totaling more than that limit, then run the GA.",
         "The GA must not assign courses to a faculty member beyond their declared unit_load ceiling."),
        ("CON-12", "Full-time faculty weekly hours capped at 40 hours",
         "Assign a full-time faculty to classes spread across multiple days, run the GA, and check the total weekly hours.",
         "The total weekly scheduled hours for a full-time faculty member must not exceed 40 hours."),
    ]),
    ("Hard Constraint - Daily Span", [
        ("CON-13", "Faculty daily teaching span must not exceed 9 hours",
         "Assign a faculty member classes that would span more than 9 hours in a single day, then run the GA.",
         "The gap from the faculty's first class start time to the last class end time on any day must not exceed 9 hours."),
        ("CON-14", "Faculty daily teaching span should meet minimum of 6 hours",
         "Assign a faculty member multiple classes in one day, run the GA, and check the total span of that day.",
         "When multiple classes are assigned on the same day, the span from the first to last class should be at least 6 hours."),
    ]),
    ("Hard Constraint - Branch Rules", [
        ("CON-15", "Faculty can teach at Main branch plus at most one other branch per week",
         "Assign a faculty to classes at three different non-Main branches across the same week, then run the GA.",
         "The GA must only allow the faculty to be scheduled at the Main branch plus one non-Main branch per week."),
        ("CON-16", "Far branch forces full-day assignment with no branch mixing",
         "Assign a faculty to a far branch (one-way travel over 4 hours) and another branch on the same day, then run the GA.",
         "The GA must not allow branch mixing on the day a faculty is at a far branch. The entire day must stay at that branch."),
        ("CON-17", "Samal branch: full-day face-to-face with online Main fill for vacant slots",
         "Assign a faculty to the Samal branch (ID 5), run the GA, and inspect the generated schedule.",
         "All Samal classes must be face-to-face. Remaining vacant slots on that day must be filled with online Main branch classes only."),
        ("CON-18", "DAPECOL branch: whole-day face-to-face with no online mode allowed",
         "Assign a faculty to the DAPECOL branch (ID 6), run the GA, and inspect the generated schedule.",
         "All DAPECOL classes must be face-to-face. Online class mode is not allowed for any DAPECOL assignment."),
    ]),
    ("Soft Constraint - Preferred Time", [
        ("CON-19", "Faculty preferred teaching time should be respected",
         "Set a faculty preferred time window (e.g., 7:00 AM to 12:00 PM), run the GA, and inspect their assigned slots.",
         "Classes assigned to that faculty should fall within the preferred time window whenever possible."),
        ("CON-20", "Faculty with no preferred time receives default scheduling",
         "Leave a faculty member's preferred time field blank, then run the GA.",
         "The faculty is scheduled using the default AM/PM balanced logic and no error or crash occurs."),
    ]),
    ("Soft Constraint - Face-to-Face Ratio", [
        ("CON-21", "At least 70% of lecture hours must be scheduled as face-to-face",
         "Run the GA with a full set of classes and check the class mode distribution in the generated output.",
         "At least 70% of total scheduled lecture contact hours should be face-to-face. Lab classes are always face-to-face and do not count toward this ratio."),
    ]),
    ("Soft Constraint - Back-to-Back Scheduling", [
        ("CON-22", "Classes within the same day should be consecutive with no large idle gaps",
         "Run the GA and inspect the daily schedule of any faculty with multiple classes on the same day.",
         "Classes for the same faculty on the same day should be scheduled back-to-back or with minimal gaps, avoiding long idle periods between sessions."),
    ]),
    ("Soft Constraint - Travel Time Between Buildings", [
        ("CON-23", "Adequate travel time gap required between rooms in different buildings",
         "Assign a faculty consecutive classes in two different buildings with non-zero travel time, then run the GA.",
         "The GA must leave a gap between those two classes that is equal to or greater than the travel time between the buildings."),
    ]),
    ("Special Rule - IAAS Institute", [
        ("CON-24", "IAAS institute lab and lecture rooms are interchangeable",
         "Create courses under the IAAS institute (ID 63), run the GA, and check which room types are assigned.",
         "For IAAS courses, the scheduler may assign either a lecture room or a lab room regardless of the course type, with no room type mismatch error."),
    ]),
]

even = False
for grp, rows in data:
    group(ws, "  " + grp)
    for r in rows:
        add_row(ws, [r[0], r[1], r[2], r[3], "", ""], even)
        even = not even

wb.save(PATH)
print("Saved:", PATH)
